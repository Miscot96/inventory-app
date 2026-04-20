import tkinter as tk
from tkinter import ttk, messagebox, simpledialog, filedialog
import sqlite3
import datetime
from openpyxl import Workbook
import tkinter.font as tkfont

from database import get_conn, init_db
from auth import hash_password, verify_password


class LoginWindow(tk.Toplevel):
    def __init__(self, master, on_success):
        super().__init__(master)
        self.title("Login")
        self.on_success = on_success
        self.resizable(False, False)

        ttk.Label(self, text="Username:").pack(padx=10, pady=(12, 2), anchor="w")
        self.username_entry = ttk.Entry(self)
        self.username_entry.pack(fill="x", padx=10)

        ttk.Label(self, text="Password:").pack(padx=10, pady=(8, 2), anchor="w")
        self.pwd_entry = ttk.Entry(self, show="*")
        self.pwd_entry.pack(fill="x", padx=10)

        btn_frame = ttk.Frame(self)
        btn_frame.pack(pady=12)
        ttk.Button(btn_frame, text="Login", command=self.try_login).grid(row=0, column=0, padx=6)
        ttk.Button(btn_frame, text="Quit", command=self.master.quit).grid(row=0, column=1, padx=6)

        self.bind("<Return>", lambda e: self.try_login())

        self.geometry("360x160")
        self.after(10, self.center_window)

    def center_window(self):
        self.update_idletasks()

        width = self.winfo_width()
        height = self.winfo_height()

        screen_width = self.winfo_screenwidth()
        screen_height = self.winfo_screenheight()

        x = (screen_width // 2) - (width // 2)
        y = (screen_height // 2) - (height // 2)

        self.geometry(f"{width}x{height}+{x}+{y}")

    def try_login(self):
        username = self.username_entry.get().strip()
        password = self.pwd_entry.get()
        if not username:
            messagebox.showwarning("Login", "Enter username", parent=self)
            return
        conn = get_conn()
        cur = conn.cursor()
        cur.execute(
            "SELECT id, salt, pwd_hash, is_admin, is_main_admin FROM users WHERE username = ?",
            (username,)
        )
        row = cur.fetchone()
        conn.close()
        if row:
            uid, salt, pwd_hash, is_admin, is_main_admin = row
            if verify_password(password, salt, pwd_hash):
                self.destroy()
                self.on_success(
                    dict(
                        id=uid,
                        username=username,
                        is_admin=bool(is_admin),
                        is_main_admin=bool(is_main_admin)
                    )
                )
                return
        messagebox.showerror("Login failed", "Invalid username or password", parent=self)


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Network Inventory Manager")
        self.geometry("1180x700")
        self.minsize(950, 600)
        self.user = None
        self.current_branch_id = None
        init_db()
        self.withdraw()
        LoginWindow(self, self.post_login)

    def post_login(self, user):
        self.user = user
        self.deiconify()
        self.create_widgets()
        self.after(10, self.maximize_window)

    def center_window(self):
        self.update_idletasks()

        width = self.winfo_width()
        height = self.winfo_height()

        screen_width = self.winfo_screenwidth()
        screen_height = self.winfo_screenheight()

        x = (screen_width // 2) - (width // 2)
        y = (screen_height // 2) - (height // 2)

        self.geometry(f"{width}x{height}+{x}+{y}")

    def center_toplevel(self, win):
        win.update_idletasks()

        width = win.winfo_width()
        height = win.winfo_height()

        screen_width = win.winfo_screenwidth()
        screen_height = win.winfo_screenheight()

        x = (screen_width // 2) - (width // 2)
        y = (screen_height // 2) - (height // 2)

        win.geometry(f"{width}x{height}+{x}+{y}")

    def maximize_window(self):
        self.update_idletasks()
        self.state("zoomed")

    def create_tree_with_scrollbars(self, parent, columns, height=None):
        outer = ttk.Frame(parent)
        outer.pack(fill="both", expand=True, padx=6, pady=6)

        tree = ttk.Treeview(outer, columns=columns, show="headings", height=height)

        y_scroll = ttk.Scrollbar(outer, orient="vertical", command=tree.yview)
        x_scroll = ttk.Scrollbar(outer, orient="horizontal", command=tree.xview)

        tree.configure(yscrollcommand=y_scroll.set, xscrollcommand=x_scroll.set)

        outer.grid_rowconfigure(0, weight=1)
        outer.grid_columnconfigure(0, weight=1)

        tree.grid(row=0, column=0, sticky="nsew")
        y_scroll.grid(row=0, column=1, sticky="ns")
        x_scroll.grid(row=1, column=0, sticky="ew")

        return tree

    def setup_columns(self, tree, specs):
        """
        specs = [(column_name, title, width, anchor, sortable_bool), ...]
        """
        if not hasattr(self, "_tree_base_headings"):
            self._tree_base_headings = {}

        tree_name = str(tree)
        self._tree_base_headings[tree_name] = {}

        for col, title, width, anchor, sortable in specs:
            self._tree_base_headings[tree_name][col] = title

            if sortable:
                tree.heading(col, text=title, command=lambda c=col: self.on_treeview_heading_click(tree, c))
            else:
                tree.heading(col, text=title)

            tree.column(col, width=width, anchor=anchor, stretch=False)

    def update_sort_arrows(self, tree, sorted_col=None, reverse=False):
        tree_name = str(tree)
        base_titles = getattr(self, "_tree_base_headings", {}).get(tree_name, {})

        for col in tree["columns"]:
            base_title = base_titles.get(col, col)

            if col == sorted_col:
                arrow = " ▼" if reverse else " ▲"
                tree.heading(col, text=base_title + arrow,
                             command=lambda c=col: self.on_treeview_heading_click(tree, c))
            else:
                tree.heading(col, text=base_title, command=lambda c=col: self.on_treeview_heading_click(tree, c))

    def autofit_column(self, tree, col):
        font = tkfont.nametofont("TkDefaultFont")

        header_text = tree.heading(col)["text"]
        max_width = font.measure(str(header_text)) + 24

        for item in tree.get_children(""):
            value = tree.set(item, col)
            cell_width = font.measure(str(value)) + 24
            if cell_width > max_width:
                max_width = cell_width

        tree.column(col, width=max(70, max_width), stretch=False)

    def bind_header_autofit(self, tree):
        def on_separator_double_click(event):
            region = tree.identify_region(event.x, event.y)
            if region != "separator":
                return

            columns = tree["columns"]
            if not columns:
                return

            current_x = 0
            for col in columns:
                current_x += tree.column(col, "width")
                if abs(event.x - current_x) <= 6:
                    self.autofit_column(tree, col)
                    break

        tree.bind("<Double-1>", on_separator_double_click, add="+")

    def format_datetime(self, value):
        if not value:
            return ""

        try:
            dt = datetime.datetime.fromisoformat(str(value))
            return dt.strftime("%Y-%m-%d %H:%M")
        except ValueError:
            return str(value).replace("T", " ")[:16]

    def create_widgets(self):
        top_outer = ttk.Frame(self)
        top_outer.pack(fill="x", padx=8, pady=6)

        # Row 1: user + branch controls
        top_row = ttk.Frame(top_outer)
        top_row.pack(fill="x", pady=(0, 6))

        left_controls = ttk.Frame(top_row)
        left_controls.pack(side="left")

        right_controls = ttk.Frame(top_row)
        right_controls.pack(side="right")

        # LEFT SIDE (branch controls)
        ttk.Label(left_controls, text="Branch:").pack(side="left")
        self.branch_combo = ttk.Combobox(left_controls, state="readonly", width=35)
        self.branch_combo.pack(side="left", padx=6)

        ttk.Button(left_controls, text="Select", command=self.select_branch).pack(side="left")
        ttk.Button(left_controls, text="Add Branch", command=self.add_branch).pack(side="left", padx=6)
        ttk.Button(left_controls, text="View Info", command=self.view_branch_info).pack(side="left", padx=6)

        if self.user.get("is_main_admin", False):
            ttk.Button(left_controls, text="Edit", command=self.edit_branch_info).pack(side="left", padx=6)
            ttk.Button(left_controls, text="Delete", command=self.delete_branch).pack(side="left", padx=6)

        ttk.Button(left_controls, text="Refresh", command=self.load_branches).pack(side="left", padx=6)

        # RIGHT SIDE (user controls)
        ttk.Label(right_controls, text=f"{self.user['username']}").pack(side="right")
        ttk.Button(right_controls, text="Password", command=self.change_my_password).pack(side="right", padx=6)

        if self.user["is_admin"]:
            ttk.Button(right_controls, text="Users", command=self.open_user_mgmt).pack(side="right", padx=6)

        # Row 2: export controls
        export_row = ttk.Frame(top_outer)
        export_row.pack(fill="x")

        ttk.Button(
            export_row,
            text="Export Materials Summary",
            command=self.export_materials_summary
        ).pack(side="left", padx=(0, 6))

        ttk.Button(
            export_row,
            text="Export Orders Summary",
            command=self.export_orders_summary
        ).pack(side="left", padx=6)

        self.load_branches()

        self.notebook = ttk.Notebook(self)
        self.notebook.pack(fill="both", expand=True, padx=8, pady=8)

        self.mat_frame = ttk.Frame(self.notebook)
        self.notebook.add(self.mat_frame, text="Materials")

        self.order_frame = ttk.Frame(self.notebook)
        self.notebook.add(self.order_frame, text="Orders")

        self.build_materials_ui()
        self.build_orders_ui()

    def change_my_password(self):
        win = tk.Toplevel(self)
        win.title("Change My Password")
        win.geometry("360x220")
        win.resizable(False, False)
        win.transient(self)
        win.grab_set()
        self.center_toplevel(win)

        ttk.Label(win, text="Current password:").pack(anchor="w", padx=10, pady=(12, 2))
        current_entry = ttk.Entry(win, show="*")
        current_entry.pack(fill="x", padx=10)

        ttk.Label(win, text="New password:").pack(anchor="w", padx=10, pady=(10, 2))
        new_entry = ttk.Entry(win, show="*")
        new_entry.pack(fill="x", padx=10)

        ttk.Label(win, text="Confirm new password:").pack(anchor="w", padx=10, pady=(10, 2))
        confirm_entry = ttk.Entry(win, show="*")
        confirm_entry.pack(fill="x", padx=10)

        def save_password():
            current_pwd = current_entry.get()
            new_pwd = new_entry.get()
            confirm_pwd = confirm_entry.get()

            if not current_pwd or not new_pwd or not confirm_pwd:
                messagebox.showwarning("Missing data", "Fill in all fields.", parent=win)
                return

            if new_pwd != confirm_pwd:
                messagebox.showerror("Mismatch", "New passwords do not match.", parent=win)
                return

            conn = get_conn()
            cur = conn.cursor()
            cur.execute(
                "SELECT salt, pwd_hash FROM users WHERE id = ?",
                (self.user["id"],)
            )
            row = cur.fetchone()

            if not row:
                conn.close()
                messagebox.showerror("Error", "User not found.", parent=win)
                return

            salt, pwd_hash = row

            if not verify_password(current_pwd, salt, pwd_hash):
                conn.close()
                messagebox.showerror("Error", "Current password is incorrect.", parent=win)
                return

            new_salt, new_hash = hash_password(new_pwd)
            cur.execute(
                "UPDATE users SET salt = ?, pwd_hash = ? WHERE id = ?",
                (new_salt, new_hash, self.user["id"])
            )
            conn.commit()
            conn.close()

            messagebox.showinfo("Success", "Password changed successfully.", parent=win)
            win.destroy()

        btn_frame = ttk.Frame(win)
        btn_frame.pack(pady=14)

        ttk.Button(btn_frame, text="Save", command=save_password).pack(side="left", padx=6)
        ttk.Button(btn_frame, text="Cancel", command=win.destroy).pack(side="left", padx=6)

    def load_branches(self):
        conn = get_conn()
        cur = conn.cursor()
        cur.execute("SELECT id, name FROM branches ORDER BY name")
        rows = cur.fetchall()
        conn.close()
        values = [f"{r[0]} - {r[1]}" for r in rows]
        self.branch_map = {f"{r[0]} - {r[1]}": r[0] for r in rows}
        self.branch_combo['values'] = values

    def select_branch(self):
        sel = self.branch_combo.get()
        if not sel:
            messagebox.showwarning("Select branch", "Choose a branch first.", parent=self)
            return

        self.current_branch_id = self.branch_map[sel]
        self.load_materials()
        self.load_orders()

    def add_branch(self):
        name = simpledialog.askstring("Add Branch", "Branch (client) name:", parent=self)
        if not name:
            return

        name = name.strip()
        if not name:
            messagebox.showwarning("Missing name", "Branch name cannot be empty.", parent=self)
            return

        desc = simpledialog.askstring("Add Branch", "Description (optional):", parent=self)

        conn = get_conn()
        cur = conn.cursor()

        cur.execute("SELECT COUNT(*) FROM branches WHERE LOWER(name) = LOWER(?)", (name,))
        exists = cur.fetchone()[0]

        if exists:
            conn.close()
            messagebox.showerror("Error", "A branch with that name already exists.", parent=self)
            return

        cur.execute(
            "INSERT INTO branches (name, description) VALUES (?, ?)",
            (name, desc if desc else None)
        )
        branch_id = cur.lastrowid
        conn.commit()
        conn.close()

        self.load_branches()

        display_value = f"{branch_id} - {name}"
        self.branch_combo.set(display_value)
        self.current_branch_id = branch_id

        self.load_materials()
        self.load_orders()

        messagebox.showinfo("Added", f"Branch '{name}' added.", parent=self)

    def edit_branch_info(self):
        if not self.user.get("is_main_admin", False):
            messagebox.showerror("Unauthorized", "Only the main admin can edit branch info.", parent=self)
            return

        sel = self.branch_combo.get()
        if not sel:
            messagebox.showwarning("Select branch", "Choose a branch first.", parent=self)
            return

        branch_id = self.branch_map[sel]

        conn = get_conn()
        cur = conn.cursor()
        cur.execute(
            "SELECT name, description FROM branches WHERE id = ?",
            (branch_id,)
        )
        row = cur.fetchone()
        conn.close()

        if not row:
            messagebox.showerror("Error", "Branch not found.", parent=self)
            return

        branch_name, description = row
        description = description or ""

        win = tk.Toplevel(self)
        win.title("Edit Branch Info")
        win.geometry("760x520")
        win.minsize(650, 420)
        win.transient(self)
        win.grab_set()
        self.center_toplevel(win)

        main = ttk.Frame(win)
        main.pack(fill="both", expand=True, padx=10, pady=10)

        ttk.Label(main, text="Branch Name:").pack(anchor="w")
        name_entry = ttk.Entry(main)
        name_entry.pack(fill="x", pady=(0, 10))
        name_entry.insert(0, branch_name)

        ttk.Label(main, text="Description / History / Important Information:").pack(anchor="w")

        text_frame = ttk.Frame(main)
        text_frame.pack(fill="both", expand=True, pady=(0, 10))

        desc_scroll = ttk.Scrollbar(text_frame, orient="vertical")
        desc_text = tk.Text(
            text_frame,
            wrap="word",
            height=18,
            yscrollcommand=desc_scroll.set
        )
        desc_scroll.config(command=desc_text.yview)

        desc_text.pack(side="left", fill="both", expand=True)
        desc_scroll.pack(side="right", fill="y")

        desc_text.insert("1.0", description)

        def save_branch_info():
            new_name = name_entry.get().strip()
            new_description = desc_text.get("1.0", "end").strip()

            if not new_name:
                messagebox.showwarning("Missing name", "Branch name cannot be empty.", parent=win)
                return

            conn = get_conn()
            cur = conn.cursor()

            cur.execute(
                "SELECT COUNT(*) FROM branches WHERE LOWER(name) = LOWER(?) AND id <> ?",
                (new_name, branch_id)
            )
            exists = cur.fetchone()[0]

            if exists:
                conn.close()
                messagebox.showerror("Error", "A different branch with that name already exists.", parent=win)
                return

            cur.execute(
                "UPDATE branches SET name = ?, description = ? WHERE id = ?",
                (new_name, new_description if new_description else None, branch_id)
            )
            conn.commit()
            conn.close()

            self.load_branches()
            updated_display = f"{branch_id} - {new_name}"
            self.branch_combo.set(updated_display)
            self.current_branch_id = branch_id
            self.load_materials()
            self.load_orders()

            messagebox.showinfo("Saved", "Branch information updated.", parent=win)
            win.destroy()

        button_frame = ttk.Frame(main)
        button_frame.pack(fill="x")

        ttk.Button(button_frame, text="Save Changes", command=save_branch_info).pack(side="left")
        ttk.Button(button_frame, text="Close", command=win.destroy).pack(side="right")

    def delete_branch(self):
        if not self.user.get("is_main_admin", False):
            messagebox.showerror(
                "Unauthorized",
                "Only the main admin can delete branches.",
                parent=self
            )
            return

        sel = self.branch_combo.get()
        if not sel:
            messagebox.showwarning("Select branch", "Choose a branch first.", parent=self)
            return

        branch_id = self.branch_map[sel]
        branch_name = sel.split(" - ", 1)[1] if " - " in sel else sel

        confirm = messagebox.askyesno(
            "Confirm Delete",
            f"Delete branch '{branch_name}'?\n\nThis will also delete all related materials and orders.",
            parent=self
        )
        if not confirm:
            return

        conn = get_conn()
        cur = conn.cursor()
        cur.execute("DELETE FROM branches WHERE id = ?", (branch_id,))
        conn.commit()
        conn.close()

        self.current_branch_id = None
        self.branch_combo.set("")
        self.load_branches()

        if hasattr(self, "mat_tree"):
            for row in self.mat_tree.get_children():
                self.mat_tree.delete(row)

        if hasattr(self, "order_tree"):
            for row in self.order_tree.get_children():
                self.order_tree.delete(row)

        messagebox.showinfo("Deleted", f"Branch '{branch_name}' deleted.", parent=self)

    def view_branch_info(self):
        sel = self.branch_combo.get()
        if not sel:
            messagebox.showwarning("Select branch", "Choose a branch first.", parent=self)
            return

        branch_id = self.branch_map[sel]

        conn = get_conn()
        cur = conn.cursor()
        cur.execute(
            "SELECT name, description FROM branches WHERE id = ?",
            (branch_id,)
        )
        row = cur.fetchone()
        conn.close()

        if not row:
            messagebox.showerror("Error", "Branch not found.", parent=self)
            return

        branch_name, description = row
        description = description or ""

        win = tk.Toplevel(self)
        win.title("Branch Info")
        win.geometry("760x520")
        win.minsize(650, 420)
        win.transient(self)
        win.grab_set()
        self.center_toplevel(win)

        main = ttk.Frame(win)
        main.pack(fill="both", expand=True, padx=10, pady=10)

        ttk.Label(main, text="Branch Name:").pack(anchor="w")
        name_entry = ttk.Entry(main)
        name_entry.pack(fill="x", pady=(0, 10))
        name_entry.insert(0, branch_name)
        name_entry.configure(state="disabled")

        ttk.Label(main, text="Description / History / Important Information:").pack(anchor="w")

        text_frame = ttk.Frame(main)
        text_frame.pack(fill="both", expand=True, pady=(0, 10))

        desc_scroll = ttk.Scrollbar(text_frame, orient="vertical")
        desc_text = tk.Text(
            text_frame,
            wrap="word",
            height=18,
            yscrollcommand=desc_scroll.set
        )
        desc_scroll.config(command=desc_text.yview)

        desc_text.pack(side="left", fill="both", expand=True)
        desc_scroll.pack(side="right", fill="y")

        desc_text.insert("1.0", description)
        desc_text.configure(state="disabled")

        button_frame = ttk.Frame(main)
        button_frame.pack(fill="x")

        ttk.Button(button_frame, text="Close", command=win.destroy).pack(side="right")

    def open_user_mgmt(self):
        if not self.user['is_admin']:
            messagebox.showerror("Unauthorized", "Only admins may manage users.", parent=self)
            return

        win = tk.Toplevel(self)
        win.title("User Management")
        win.geometry("760x420")
        win.resizable(False, False)
        win.transient(self)

        frm = ttk.Frame(win)
        frm.pack(fill="both", expand=True, padx=8, pady=8)

        columns = ("id", "username", "is_admin", "created_at")
        tree = self.create_tree_with_scrollbars(frm, columns)

        specs = [
            ("id", "ID", 70, "center", True),
            ("username", "Username", 180, "w", True),
            ("is_admin", "Admin", 100, "center", True),
            ("created_at", "Created At", 300, "w", True),
        ]
        self.setup_columns(tree, specs)
        self.bind_header_autofit(tree)
        self.center_toplevel(win)

        def load_users():
            for r in tree.get_children():
                tree.delete(r)

            conn = get_conn()
            cur = conn.cursor()
            cur.execute("SELECT id, username, is_admin, created_at FROM users ORDER BY id")

            for user_id, username, is_admin, created_at in cur.fetchall():
                admin_text = "Yes" if is_admin else "No"
                tree.insert("", "end", values=(user_id, username, admin_text, self.format_datetime(created_at)))

            conn.close()

        def add_user():
            uname = simpledialog.askstring("New user", "Username:", parent=win)
            if not uname:
                return

            pwd = simpledialog.askstring("New user", f"Password for {uname}:", show="*", parent=win)
            if pwd is None:
                return

            is_admin = messagebox.askyesno("Is Admin", "Should this user be admin?", parent=win)
            salt, pwd_hash = hash_password(pwd)

            conn = get_conn()
            cur = conn.cursor()
            try:
                cur.execute(
                    "INSERT INTO users (username, salt, pwd_hash, is_admin, created_at) VALUES (?, ?, ?, ?, ?)",
                    (uname, salt, pwd_hash, 1 if is_admin else 0, datetime.datetime.utcnow().isoformat())
                )
                conn.commit()
                messagebox.showinfo("Added", f"User {uname} added.", parent=win)
            except sqlite3.IntegrityError:
                messagebox.showerror("Error", "Username already exists.", parent=win)
            conn.close()

            load_users()

        def delete_user():
            sel = tree.selection()
            if not sel:
                messagebox.showwarning("Select", "Select a user to delete.", parent=win)
                return

            uid = tree.item(sel[0])['values'][0]

            if uid == self.user['id']:
                messagebox.showerror("Error", "You cannot delete yourself while logged in.", parent=win)
                return

            if messagebox.askyesno("Confirm", "Delete selected user?", parent=win):
                conn = get_conn()
                cur = conn.cursor()
                cur.execute("DELETE FROM users WHERE id = ?", (uid,))
                conn.commit()
                conn.close()
                load_users()

        def reset_user_password():
            sel = tree.selection()
            if not sel:
                messagebox.showwarning("Select", "Select a user first.", parent=win)
                return

            values = tree.item(sel[0])["values"]
            target_user_id = values[0]
            target_username = values[1]

            conn = get_conn()
            cur = conn.cursor()
            cur.execute(
                "SELECT is_admin, is_main_admin FROM users WHERE id = ?",
                (target_user_id,)
            )
            row = cur.fetchone()

            if not row:
                conn.close()
                messagebox.showerror("Error", "Selected user not found.", parent=win)
                return

            target_is_admin, _ = row

            # You must change your own password through the personal password window
            if target_user_id == self.user["id"]:
                conn.close()
                messagebox.showwarning(
                    "Not allowed",
                    "Use 'Change My Password' to change your own password.",
                    parent=win
                )
                return

            # Normal admins cannot reset another admin or main admin
            if not self.user.get("is_main_admin", False) and target_is_admin:
                conn.close()
                messagebox.showerror(
                    "Not allowed",
                    "Only the main admin can reset another admin's password.",
                    parent=win
                )
                return

            new_pwd = simpledialog.askstring(
                "Reset Password",
                f"Enter a new password for {target_username}:",
                show="*",
                parent=win
            )
            if not new_pwd:
                conn.close()
                return

            confirm_pwd = simpledialog.askstring(
                "Reset Password",
                f"Confirm new password for {target_username}:",
                show="*",
                parent=win
            )
            if not confirm_pwd:
                conn.close()
                return

            if new_pwd != confirm_pwd:
                conn.close()
                messagebox.showerror("Mismatch", "Passwords do not match.", parent=win)
                return

            new_salt, new_hash = hash_password(new_pwd)

            cur.execute(
                "UPDATE users SET salt = ?, pwd_hash = ? WHERE id = ?",
                (new_salt, new_hash, target_user_id)
            )
            conn.commit()
            conn.close()

            messagebox.showinfo("Success", f"Password for {target_username} was updated.", parent=win)

        buttons_frame = ttk.Frame(frm)
        buttons_frame.pack(fill="x", pady=6)

        ttk.Button(buttons_frame, text="Add User", command=add_user).pack(side="left")
        ttk.Button(buttons_frame, text="Delete User", command=delete_user).pack(side="left", padx=6)
        ttk.Button(buttons_frame, text="Reset Password", command=reset_user_password).pack(side="left", padx=6)
        ttk.Button(buttons_frame, text="Refresh", command=load_users).pack(side="left", padx=6)

        load_users()

    def build_materials_ui(self):
        top = ttk.Frame(self.mat_frame)
        top.pack(fill="x", pady=6)

        ttk.Button(top, text="Add Material Reference", command=self.add_material_reference).pack(side="left")
        ttk.Button(top, text="Refresh Materials", command=self.load_materials).pack(side="left", padx=6)

        search_frame = ttk.Frame(self.mat_frame)
        search_frame.pack(fill="x", padx=6, pady=(0, 6))

        ttk.Label(search_frame, text="Search:").pack(side="left")
        self.material_search_entry = ttk.Entry(search_frame)
        self.material_search_entry.pack(side="left", fill="x", expand=True, padx=6)
        self.material_search_entry.bind("<Return>", lambda e: self.load_materials())

        ttk.Button(search_frame, text="Search", command=self.load_materials).pack(side="left", padx=4)
        ttk.Button(search_frame, text="Clear", command=self.clear_material_search).pack(side="left")

        cols = ("id", "reference", "color", "available", "needed", "missing", "orders", "unit")
        tree = self.create_tree_with_scrollbars(self.mat_frame, cols)

        specs = [
            ("id", "ID", 70, "center", True),
            ("reference", "Reference", 220, "w", True),
            ("color", "Color", 120, "w", True),
            ("available", "Available", 90, "center", True),
            ("needed", "Needed", 90, "center", True),
            ("missing", "Missing", 90, "center", True),
            ("orders", "Orders", 80, "center", True),
            ("unit", "Unit", 70, "center", True),
        ]
        self.setup_columns(tree, specs)
        self.bind_header_autofit(tree)

        self.mat_tree = tree
        tree.tag_configure("missing", background="#ffe5e5")
        tree.tag_configure("ok", background="#e8f5e9")

        menu = tk.Menu(self, tearoff=0)
        menu.add_command(label="Edit Material", command=self.edit_material)
        menu.add_command(label="Show Orders Using This Material", command=self.show_material_orders)
        menu.add_command(label="Delete Material Reference", command=self.delete_material_reference)

        def show_menu(event):
            try:
                row_id = tree.identify_row(event.y)
                if row_id:
                    tree.selection_set(row_id)
                menu.tk_popup(event.x_root, event.y_root)
            finally:
                menu.grab_release()

        tree.bind("<Button-3>", show_menu)

    def add_material_reference(self):
        if not self.current_branch_id:
            messagebox.showwarning("Choose branch", "Select branch first.", parent=self)
            return

        win = tk.Toplevel(self)
        win.title("Add Material Reference")
        win.geometry("430x380")
        win.resizable(False, False)
        win.transient(self)
        win.grab_set()
        self.center_toplevel(win)

        ttk.Label(win, text="Reference code/name:").pack(anchor="w", padx=10, pady=(12, 2))
        ref_entry = ttk.Entry(win)
        ref_entry.pack(fill="x", padx=10)

        ttk.Label(win, text="Description (optional):").pack(anchor="w", padx=10, pady=(10, 2))
        desc_entry = ttk.Entry(win)
        desc_entry.pack(fill="x", padx=10)

        ttk.Label(win, text="First color (optional):").pack(anchor="w", padx=10, pady=(10, 2))
        color_entry = ttk.Entry(win)
        color_entry.pack(fill="x", padx=10)

        ttk.Label(win, text="Unit:").pack(anchor="w", padx=10, pady=(10, 2))
        unit_combo = ttk.Combobox(win, values=["m", "pcs"], state="readonly")
        unit_combo.pack(fill="x", padx=10)
        unit_combo.set("m")

        ttk.Label(win, text="Available quantity (only if color is entered):").pack(anchor="w", padx=10, pady=(10, 2))
        qty_entry = ttk.Entry(win)
        qty_entry.pack(fill="x", padx=10)

        def save_material():
            ref = ref_entry.get().strip()
            desc = desc_entry.get().strip()
            color = color_entry.get().strip()
            unit = unit_combo.get().strip()
            qty_text = qty_entry.get().strip()

            if not ref:
                messagebox.showwarning("Missing reference", "Enter a reference code/name.", parent=win)
                return

            conn = get_conn()
            cur = conn.cursor()

            try:
                cur.execute(
                    "INSERT INTO materials (branch_id, reference, description) VALUES (?, ?, ?)",
                    (self.current_branch_id, ref, desc if desc else None)
                )
                material_id = cur.lastrowid

                if color:
                    if not qty_text:
                        messagebox.showwarning("Missing quantity", "Enter an available quantity for the color.",
                                               parent=win)
                        conn.close()
                        return

                    try:
                        avail = float(qty_text)
                    except ValueError:
                        messagebox.showwarning("Invalid quantity", "Available quantity must be a number.", parent=win)
                        conn.close()
                        return

                    cur.execute(
                        "INSERT INTO material_colors (material_id, color, unit, available_qty) VALUES (?, ?, ?, ?)",
                        (material_id, color, unit, avail)
                    )

                conn.commit()
                messagebox.showinfo("Added", "Material reference added successfully.", parent=win)
                win.destroy()
                self.load_materials()

            except sqlite3.IntegrityError:
                messagebox.showerror("Exists", "This reference already exists for this branch.", parent=win)
                conn.close()
                return

            conn.close()

        btn_frame = ttk.Frame(win)
        btn_frame.pack(pady=16)

        ttk.Button(btn_frame, text="Save", command=save_material).pack(side="left", padx=6)
        ttk.Button(btn_frame, text="Cancel", command=win.destroy).pack(side="left", padx=6)

    def edit_material(self):
        sel = self.mat_tree.selection()
        if not sel:
            messagebox.showwarning("Select", "Select a material row first.", parent=self)
            return

        values = self.mat_tree.item(sel[0])["values"]
        row_id = values[0]
        reference = values[1]
        color = values[2]
        unit = values[7]
        available = values[3]

        conn = get_conn()
        cur = conn.cursor()

        if color == "-":
            material_id = row_id
            cur.execute(
                "SELECT description FROM materials WHERE id = ?",
                (material_id,)
            )
            row = cur.fetchone()
            description = row[0] if row else ""
            material_color_id = None
        else:
            material_color_id = row_id
            cur.execute("""
                SELECT m.id, m.description, mc.color, mc.unit, mc.available_qty
                FROM material_colors mc
                JOIN materials m ON mc.material_id = m.id
                WHERE mc.id = ?
            """, (material_color_id,))
            row = cur.fetchone()
            if not row:
                conn.close()
                messagebox.showerror("Error", "Material row not found.", parent=self)
                return
            material_id, description, color, unit, available = row

        conn.close()

        win = tk.Toplevel(self)
        win.title("Edit Material")
        win.geometry("430x380")
        win.resizable(False, False)
        win.transient(self)
        win.grab_set()
        self.center_toplevel(win)

        ttk.Label(win, text="Reference code/name:").pack(anchor="w", padx=10, pady=(12, 2))
        ref_entry = ttk.Entry(win)
        ref_entry.pack(fill="x", padx=10)
        ref_entry.insert(0, reference)

        ttk.Label(win, text="Description (optional):").pack(anchor="w", padx=10, pady=(10, 2))
        desc_entry = ttk.Entry(win)
        desc_entry.pack(fill="x", padx=10)
        desc_entry.insert(0, description or "")

        ttk.Label(win, text="Color (optional):").pack(anchor="w", padx=10, pady=(10, 2))
        color_entry = ttk.Entry(win)
        color_entry.pack(fill="x", padx=10)
        if color != "-":
            color_entry.insert(0, color)

        ttk.Label(win, text="Unit:").pack(anchor="w", padx=10, pady=(10, 2))
        unit_combo = ttk.Combobox(win, values=["m", "pcs"], state="readonly")
        unit_combo.pack(fill="x", padx=10)
        unit_combo.set(unit if unit and unit != "-" else "m")

        ttk.Label(win, text="Available quantity:").pack(anchor="w", padx=10, pady=(10, 2))
        qty_entry = ttk.Entry(win)
        qty_entry.pack(fill="x", padx=10)
        if color != "-":
            qty_entry.insert(0, str(available))

        def save_changes():
            new_ref = ref_entry.get().strip()
            new_desc = desc_entry.get().strip()
            new_color = color_entry.get().strip()
            new_unit = unit_combo.get().strip()
            qty_text = qty_entry.get().strip()

            if not new_ref:
                messagebox.showwarning("Missing reference", "Enter a reference code/name.", parent=win)
                return

            conn = get_conn()
            cur = conn.cursor()

            try:
                cur.execute(
                    "UPDATE materials SET reference = ?, description = ? WHERE id = ?",
                    (new_ref, new_desc if new_desc else None, material_id)
                )

                if material_color_id is not None:
                    if not new_color:
                        messagebox.showwarning("Missing color", "Color cannot be empty for an existing color row.",
                                               parent=win)
                        conn.close()
                        return

                    try:
                        new_available = float(qty_text)
                    except ValueError:
                        messagebox.showwarning("Invalid quantity", "Available quantity must be a number.", parent=win)
                        conn.close()
                        return

                    cur.execute(
                        "UPDATE material_colors SET color = ?, unit = ?, available_qty = ? WHERE id = ?",
                        (new_color, new_unit, new_available, material_color_id)
                    )
                else:
                    # reference-only row can be turned into a color row if user fills color
                    if new_color:
                        try:
                            new_available = float(qty_text)
                        except ValueError:
                            messagebox.showwarning("Invalid quantity", "Available quantity must be a number.",
                                                   parent=win)
                            conn.close()
                            return

                        cur.execute(
                            "INSERT INTO material_colors (material_id, color, unit, available_qty) VALUES (?, ?, ?, ?)",
                            (material_id, new_color, new_unit, new_available)
                        )

                conn.commit()
                conn.close()
                messagebox.showinfo("Saved", "Material updated.", parent=win)
                win.destroy()
                self.load_materials()

            except sqlite3.IntegrityError:
                conn.close()
                messagebox.showerror("Error", "That reference/color combination already exists.", parent=win)

        btn_frame = ttk.Frame(win)
        btn_frame.pack(pady=16)

        ttk.Button(btn_frame, text="Save", command=save_changes).pack(side="left", padx=6)
        ttk.Button(btn_frame, text="Cancel", command=win.destroy).pack(side="left", padx=6)

    def add_color_variant_to_selected_reference(self):
        sel = self.mat_tree.selection()
        if not sel:
            messagebox.showwarning("Select", "Select a material reference or color row first.", parent=self)
            return

        values = self.mat_tree.item(sel[0])["values"]
        row_id = values[0]
        color = values[2]

        conn = get_conn()
        cur = conn.cursor()

        if color == "-":
            material_id = row_id
            cur.execute("SELECT reference FROM materials WHERE id = ?", (material_id,))
            row = cur.fetchone()
            reference = row[0] if row else ""
        else:
            material_color_id = row_id
            cur.execute("""
                SELECT m.id, m.reference
                FROM material_colors mc
                JOIN materials m ON mc.material_id = m.id
                WHERE mc.id = ?
            """, (material_color_id,))
            row = cur.fetchone()
            if not row:
                conn.close()
                messagebox.showerror("Error", "Material row not found.", parent=self)
                return
            material_id, reference = row

        conn.close()

        win = tk.Toplevel(self)
        win.title(f"Add Color Variant to {reference}")
        win.geometry("420x240")
        win.resizable(False, False)
        win.transient(self)
        win.grab_set()
        self.center_toplevel(win)

        ttk.Label(win, text="Color:").pack(anchor="w", padx=10, pady=(12, 2))
        color_entry = ttk.Entry(win)
        color_entry.pack(fill="x", padx=10)

        ttk.Label(win, text="Unit:").pack(anchor="w", padx=10, pady=(10, 2))
        unit_combo = ttk.Combobox(win, values=["m", "pcs"], state="readonly")
        unit_combo.pack(fill="x", padx=10)
        unit_combo.set("m")

        ttk.Label(win, text="Available quantity:").pack(anchor="w", padx=10, pady=(10, 2))
        qty_entry = ttk.Entry(win)
        qty_entry.pack(fill="x", padx=10)

        def save_variant():
            new_color = color_entry.get().strip()
            new_unit = unit_combo.get().strip()
            qty_text = qty_entry.get().strip()

            if not new_color:
                messagebox.showwarning("Missing color", "Enter a color name/code.", parent=win)
                return

            try:
                new_available = float(qty_text)
            except ValueError:
                messagebox.showwarning("Invalid quantity", "Available quantity must be a number.", parent=win)
                return

            conn = get_conn()
            cur = conn.cursor()
            try:
                cur.execute(
                    "INSERT INTO material_colors (material_id, color, unit, available_qty) VALUES (?, ?, ?, ?)",
                    (material_id, new_color, new_unit, new_available)
                )
                conn.commit()
                conn.close()
                messagebox.showinfo("Added", "Color variant added.", parent=win)
                win.destroy()
                self.load_materials()
            except sqlite3.IntegrityError:
                conn.close()
                messagebox.showerror("Exists", "This color already exists for that reference.", parent=win)

        btn_frame = ttk.Frame(win)
        btn_frame.pack(pady=16)

        ttk.Button(btn_frame, text="Save", command=save_variant).pack(side="left", padx=6)
        ttk.Button(btn_frame, text="Cancel", command=win.destroy).pack(side="left", padx=6)

    def clear_material_search(self):
        self.material_search_entry.delete(0, tk.END)
        self.load_materials()

    def get_sort_value(self, col, value):
        numeric_columns = {
            "id", "available", "needed", "missing", "orders", "order_id",
            "needed_qty", "output_qty", "qty_per_piece"
        }

        if col in numeric_columns:
            try:
                return float(value)
            except (ValueError, TypeError):
                return 0.0

        return str(value).strip().lower()

    def is_treeview_column_ascending(self, tree, col):
        values = [self.get_sort_value(col, tree.set(item, col)) for item in tree.get_children("")]

        if len(values) < 2:
            return True

        return values == sorted(values)

    def on_treeview_heading_click(self, tree, col):
        tree_name = str(tree)

        if not hasattr(self, "_sort_states"):
            self._sort_states = {}

        last_col, last_reverse = self._sort_states.get(tree_name, (None, None))

        if last_col == col:
            reverse = not last_reverse
        else:
            reverse = self.is_treeview_column_ascending(tree, col)

        self.sort_treeview(tree, col, reverse)
        self._sort_states[tree_name] = (col, reverse)
        self.update_sort_arrows(tree, col, reverse)

    def sort_treeview(self, tree, col, reverse):
        data = [(tree.set(item, col), item) for item in tree.get_children("")]

        data.sort(key=lambda item: self.get_sort_value(col, item[0]), reverse=reverse)

        for index, (_, item) in enumerate(data):
            tree.move(item, "", index)

    def load_materials(self):
        for r in self.mat_tree.get_children():
            self.mat_tree.delete(r)

        if not self.current_branch_id:
            return

        search_text = self.material_search_entry.get().strip().lower()

        conn = get_conn()
        cur = conn.cursor()

        cur.execute("""
            SELECT
                m.id,
                m.reference,
                mc.id,
                mc.color,
                mc.available_qty,
                mc.unit
            FROM materials m
            LEFT JOIN material_colors mc ON mc.material_id = m.id
            WHERE m.branch_id = ?
            ORDER BY m.reference, mc.color
        """, (self.current_branch_id,))

        rows = cur.fetchall()

        for material_id, reference, mc_id, color, available_qty, unit in rows:
            reference_text = (reference or "").lower()
            color_text = (color or "").lower()

            if search_text and search_text not in reference_text and search_text not in color_text:
                continue

            if mc_id is None:
                self.mat_tree.insert("", "end", values=(
                    material_id,
                    reference,
                    "-",
                    0.0,
                    0.0,
                    0.0,
                    0,
                    "-"
                ), tags=("ok",))
            else:
                cur.execute("""
                    SELECT IFNULL(SUM(oi.needed_qty), 0), COUNT(DISTINCT oi.order_id)
                    FROM order_items oi
                    JOIN orders o ON oi.order_id = o.id
                    WHERE oi.material_color_id = ? AND o.branch_id = ?
                """, (mc_id, self.current_branch_id))

                needed, order_count = cur.fetchone()
                needed = float(needed or 0)
                available = float(available_qty or 0)
                missing = max(0.0, needed - available)
                row_tag = "missing" if missing > 0 else "ok"

                self.mat_tree.insert("", "end", values=(
                    mc_id,
                    reference,
                    color,
                    available,
                    needed,
                    missing,
                    int(order_count or 0),
                    unit
                ), tags=(row_tag,))

        self.update_sort_arrows(self.mat_tree)

        conn.close()

    def edit_available_qty(self):
        sel = self.mat_tree.selection()
        if not sel:
            messagebox.showwarning("Select", "Select a material color row.", parent=self)
            return

        values = self.mat_tree.item(sel[0])["values"]
        mc_id = values[0]
        color_value = values[2]

        if color_value == "-":
            messagebox.showwarning("Not allowed", "This reference has no color variant yet.", parent=self)
            return

        newq = simpledialog.askstring(
            "Available quantity",
            "New available quantity (numeric):",
            parent=self
        )
        if newq is None or newq.strip() == "":
            return

        try:
            newqf = float(newq)
        except (ValueError, TypeError):
            messagebox.showerror("Invalid", "Enter valid numeric value.", parent=self)
            return

        if newqf < 0:
            messagebox.showerror("Invalid", "Quantity cannot be negative.", parent=self)
            return

        conn = get_conn()
        cur = conn.cursor()
        cur.execute(
            "UPDATE material_colors SET available_qty = ? WHERE id = ?",
            (round(newqf, 2), mc_id)
        )
        conn.commit()
        conn.close()
        self.load_materials()

    def delete_material_color(self):
        sel = self.mat_tree.selection()
        if not sel:
            messagebox.showwarning("Select", "Select a material color row.", parent=self)
            return

        values = self.mat_tree.item(sel[0])["values"]
        mc_id = values[0]
        color_value = values[2]

        if color_value == "-":
            messagebox.showwarning("Not allowed", "This reference has no color variant to delete.", parent=self)
            return

        if messagebox.askyesno("Confirm", "Delete this color variant?", parent=self):
            conn = get_conn()
            cur = conn.cursor()
            cur.execute("DELETE FROM material_colors WHERE id = ?", (mc_id,))
            conn.commit()
            conn.close()
            self.load_materials()

    def delete_material_reference(self):
        sel = self.mat_tree.selection()
        if not sel:
            messagebox.showwarning("Select", "Select a reference row to remove.", parent=self)
            return

        values = self.mat_tree.item(sel[0])["values"]
        row_id = values[0]
        color_value = values[2]

        conn = get_conn()
        cur = conn.cursor()

        if color_value == "-":
            mat_id = row_id
        else:
            cur.execute(
                "SELECT material_id FROM material_colors WHERE id = ?",
                (row_id,)
            )
            r = cur.fetchone()
            if not r:
                conn.close()
                return
            mat_id = r[0]

        conn.close()

        if messagebox.askyesno("Confirm", "Delete entire material reference and all its color variants?", parent=self):
            conn = get_conn()
            cur = conn.cursor()
            cur.execute("DELETE FROM materials WHERE id = ?", (mat_id,))
            conn.commit()
            conn.close()
            self.load_materials()

    def show_material_orders(self):
        sel = self.mat_tree.selection()
        if not sel:
            messagebox.showwarning("Select", "Select a material color row first.", parent=self)
            return

        values = self.mat_tree.item(sel[0])["values"]
        row_id = values[0]
        reference = values[1]
        color = values[2]

        if color == "-":
            messagebox.showwarning("Not allowed", "This reference has no color variant yet.", parent=self)
            return

        win = tk.Toplevel(self)
        win.title(f"Orders Using {reference} - {color}")
        win.geometry("920x420")
        win.resizable(False, False)
        win.transient(self)
        win.grab_set()
        self.center_toplevel(win)

        cols = ("order_id", "order_name", "product_name", "output_qty", "qty_per_piece", "needed_qty", "unit", "created_at")
        tree = self.create_tree_with_scrollbars(win, cols)

        specs = [
            ("order_id", "Order ID", 80, "center", True),
            ("order_name", "Order Name", 160, "w", True),
            ("product_name", "Produced Item", 160, "w", True),
            ("output_qty", "Output Qty", 90, "center", True),
            ("qty_per_piece", "Qty / Output", 100, "center", True),
            ("needed_qty", "Total Needed", 110, "center", True),
            ("unit", "Unit", 70, "center", True),
            ("created_at", "Created At", 170, "w", True),
        ]
        self.setup_columns(tree, specs)
        self.bind_header_autofit(tree)

        conn = get_conn()
        cur = conn.cursor()
        cur.execute("""
            SELECT
                o.id,
                o.name,
                o.product_name,
                o.output_qty,
                mc.unit,
                oi.qty_per_piece,
                oi.needed_qty,
                o.created_at
            FROM order_items oi
            JOIN orders o ON oi.order_id = o.id
            JOIN material_colors mc ON oi.material_color_id = mc.id
            WHERE oi.material_color_id = ? AND o.branch_id = ?
            ORDER BY o.created_at DESC
        """, (row_id, self.current_branch_id))

        for r in cur.fetchall():
            tree.insert("", "end", values=(
                r[0],
                r[1] or "",
                r[2] or "",
                float(r[3] or 0),
                float(r[5] or 0),
                float(r[6] or 0),
                r[4],
                self.format_datetime(r[7]),
            ))

        conn.close()

    def build_orders_ui(self):
        top = ttk.Frame(self.order_frame)
        top.pack(fill="x", pady=6)

        ttk.Button(top, text="Create Order", command=self.create_order).pack(side="left", padx=6)
        ttk.Button(top, text="Refresh Orders", command=self.load_orders).pack(side="left", padx=6)
        ttk.Button(top, text="Show Materials Summary", command=self.show_materials_summary_report).pack(side="left", padx=6)

        cols = ("order_id", "name", "product_name", "output_qty", "output_unit", "created_by", "created_at", "notes")
        tree = self.create_tree_with_scrollbars(self.order_frame, cols)

        specs = [
            ("order_id", "Order ID", 80, "center", True),
            ("name", "Order Name", 170, "w", True),
            ("product_name", "Produced Item", 170, "w", True),
            ("output_qty", "Production Quantity", 90, "center", True),
            ("output_unit", "Unit", 70, "center", True),
            ("created_by", "Created By", 110, "w", True),
            ("created_at", "Created At", 170, "w", True),
            ("notes", "Notes", 220, "w", True),
        ]
        self.setup_columns(tree, specs)
        self.bind_header_autofit(tree)

        tree.bind("<Double-1>", self.open_order_items)
        self.order_tree = tree
        self.bind_header_autofit(tree)

        menu = tk.Menu(self, tearoff=0)
        menu.add_command(label="Edit Order", command=self.edit_selected_order)
        menu.add_command(label="Delete Order", command=self.delete_selected_order)

        def show_menu(event):
            try:
                row_id = tree.identify_row(event.y)
                if row_id:
                    tree.selection_set(row_id)
                menu.tk_popup(event.x_root, event.y_root)
            finally:
                menu.grab_release()

        tree.bind("<Button-3>", show_menu)

    def delete_selected_order(self):
        sel = self.order_tree.selection()
        if not sel:
            messagebox.showwarning("Select", "Select an order first.", parent=self)
            return

        order_id = self.order_tree.item(sel[0])["values"][0]

        if messagebox.askyesno("Confirm", f"Delete order #{order_id}?", parent=self):
            conn = get_conn()
            cur = conn.cursor()
            cur.execute("DELETE FROM orders WHERE id = ?", (order_id,))
            conn.commit()
            conn.close()

            self.load_orders()
            self.load_materials()

    def load_orders(self):
        for r in self.order_tree.get_children():
            self.order_tree.delete(r)

        if not self.current_branch_id:
            return

        conn = get_conn()
        cur = conn.cursor()
        cur.execute("""
            SELECT
                o.id,
                o.name,
                o.product_name,
                o.output_qty,
                o.output_unit,
                u.username,
                o.created_at,
                o.notes
            FROM orders o
            LEFT JOIN users u ON o.created_by = u.id
            WHERE o.branch_id = ?
            ORDER BY o.created_at DESC
        """, (self.current_branch_id,))

        for r in cur.fetchall():
            self.order_tree.insert("", "end", values=(
                r[0],
                r[1] or "",
                r[2] or "",
                float(r[3] or 0),
                r[4],
                r[5] or "",
                self.format_datetime(r[6]),
                r[7] or ""
            ))

        self.update_sort_arrows(self.order_tree)

        conn.close()

    def create_order(self):
        self.open_order_editor()

    def edit_selected_order(self):
        sel = self.order_tree.selection()
        if not sel:
            messagebox.showwarning("Select", "Select an order first.", parent=self)
            return

        order_id = self.order_tree.item(sel[0])["values"][0]
        self.open_order_editor(order_id)

    def open_order_editor(self, order_id=None):
        if not self.current_branch_id:
            messagebox.showwarning("Choose branch", "Select branch first.", parent=self)
            return

        is_edit = order_id is not None


        win = tk.Toplevel(self)
        win.title("Edit Order" if is_edit else "Create Order")
        win.geometry("980x650")
        win.transient(self)
        win.grab_set()
        self.center_toplevel(win)

        top_frame = ttk.Frame(win)
        top_frame.pack(fill="x", padx=8, pady=8)

        ttk.Label(top_frame, text="Order name:").grid(row=0, column=0, sticky="w", padx=(0, 6))
        name_entry = ttk.Entry(top_frame, width=30)
        name_entry.grid(row=0, column=1, sticky="w")

        ttk.Label(top_frame, text="Produced item:").grid(row=0, column=2, sticky="w", padx=(12, 6))
        product_entry = ttk.Entry(top_frame, width=30)
        product_entry.grid(row=0, column=3, sticky="w")

        ttk.Label(top_frame, text="Production Quantity:").grid(row=1, column=0, sticky="w", padx=(0, 6), pady=(10, 0))
        output_qty_entry = ttk.Entry(top_frame, width=18)
        output_qty_entry.grid(row=1, column=1, sticky="w", pady=(10, 0))
        output_qty_entry.insert(0, "1")

        ttk.Label(top_frame, text="Unit:").grid(row=1, column=2, sticky="w", padx=(12, 6), pady=(10, 0))
        output_unit_combo = ttk.Combobox(top_frame, values=["pcs", "m"], state="readonly", width=10)
        output_unit_combo.grid(row=1, column=3, sticky="w", pady=(10, 0))
        output_unit_combo.set("pcs")

        ttk.Label(top_frame, text="Notes:").grid(row=2, column=0, sticky="w", padx=(0, 6), pady=(10, 0))
        notes_entry = ttk.Entry(top_frame, width=70)
        notes_entry.grid(row=2, column=1, columnspan=3, sticky="w", pady=(10, 0))

        search_frame = ttk.Frame(win)
        search_frame.pack(fill="x", padx=8, pady=(0, 8))

        ttk.Label(search_frame, text="Search material:").pack(side="left")
        search_entry = ttk.Entry(search_frame)
        search_entry.pack(side="left", fill="x", expand=True, padx=6)

        cols = (
        "mc_id", "reference", "color", "available_qty", "needed_qty", "missing", "orders", "qty_per_piece", "unit")
        tree = self.create_tree_with_scrollbars(win, cols)

        specs = [
            ("mc_id", "ID", 70, "center", True),
            ("reference", "Reference", 220, "w", True),
            ("color", "Color", 110, "w", True),
            ("available_qty", "Available", 90, "center", True),
            ("needed_qty", "Total Needed", 100, "center", True),
            ("missing", "Missing", 90, "center", True),
            ("orders", "Orders", 80, "center", True),
            ("qty_per_piece", "Material per Unit", 130, "center", True),
            ("unit", "Unit", 70, "center", True),
        ]
        self.setup_columns(tree, specs)
        self.bind_header_autofit(tree)

        conn = get_conn()
        cur = conn.cursor()
        cur.execute("""
            SELECT
                mc.id,
                m.reference,
                mc.color,
                mc.available_qty,
                IFNULL(SUM(oi.needed_qty), 0) AS total_needed_all_orders,
                COUNT(DISTINCT oi.order_id) AS orders_count,
                mc.unit
            FROM material_colors mc
            JOIN materials m ON mc.material_id = m.id
            LEFT JOIN order_items oi ON oi.material_color_id = mc.id
            LEFT JOIN orders o ON oi.order_id = o.id AND o.branch_id = m.branch_id
            WHERE m.branch_id = ?
            GROUP BY mc.id, m.reference, mc.color, mc.available_qty, mc.unit
            ORDER BY m.reference, mc.color
        """, (self.current_branch_id,))
        all_rows = cur.fetchall()

        existing_items = {}
        if is_edit:
            cur.execute("""
                SELECT name, product_name, output_qty, output_unit, notes
                FROM orders
                WHERE id = ?
            """, (order_id,))
            order_row = cur.fetchone()
            if order_row:
                name_entry.insert(0, order_row[0] or "")
                product_entry.insert(0, order_row[1] or "")
                output_qty_entry.delete(0, tk.END)
                output_qty_entry.insert(0, str(order_row[2] or 1))
                output_unit_combo.set(order_row[3] or "pcs")
                notes_entry.insert(0, order_row[4] or "")

            cur.execute("""
                SELECT material_color_id, qty_per_piece, needed_qty
                FROM order_items
                WHERE order_id = ?
            """, (order_id,))
            for mc_id, qty_per_piece, needed_qty in cur.fetchall():
                existing_items[mc_id] = (float(qty_per_piece or 0), float(needed_qty or 0))

        conn.close()

        def get_output_qty():
            try:
                val = float(output_qty_entry.get().strip())
                return val if val > 0 else 1.0
            except ValueError:
                return 1.0

        def populate_tree(filter_text=""):
            for item in tree.get_children():
                tree.delete(item)

            filter_text = filter_text.strip().lower()
            output_qty = get_output_qty()

            for mc_id, reference, color, available_qty, total_needed_all_orders, orders_count, unit in all_rows:
                ref_text = (reference or "").lower()
                color_text = (color or "").lower()

                if filter_text and filter_text not in ref_text and filter_text not in color_text:
                    continue

                qty_per_output = 0.0
                total_needed_this_order = 0.0

                if mc_id in existing_items:
                    qty_per_output, total_needed_this_order = existing_items[mc_id]

                if not is_edit:
                    total_needed_this_order = output_qty * qty_per_output

                missing = max(0.0, total_needed_this_order - float(available_qty or 0))

                tree.insert("", "end", values=(
                    mc_id,
                    reference,
                    color,
                    float(available_qty or 0),
                    round(total_needed_this_order, 2),
                    round(missing, 2),
                    int(orders_count or 0),
                    round(qty_per_output, 2),
                    unit
                ))

        def update_search(event=None):
            populate_tree(search_entry.get())

        def recalculate_totals(event=None):
            output_qty = get_output_qty()

            for item in tree.get_children():
                vals = list(tree.item(item, "values"))
                available_qty = float(vals[3])
                qty_per_output = float(vals[7])
                total_needed = round(output_qty * qty_per_output, 2)
                missing = round(max(0.0, total_needed - available_qty), 2)
                vals[4] = total_needed
                vals[5] = missing
                tree.item(item, values=vals)

        search_entry.bind("<KeyRelease>", update_search)
        search_entry.bind("<Return>", update_search)
        output_qty_entry.bind("<KeyRelease>", recalculate_totals)

        populate_tree()

        def set_qty_for_selected():
            sel = tree.selection()
            if not sel:
                messagebox.showwarning("Select", "Select rows to set quantity per output.", parent=win)
                return

            val = simpledialog.askstring("Qty / Output", "Material quantity per output unit (numeric):", parent=win)
            if val is None or val.strip() == "":
                return

            try:
                qty_per_output = float(val)
            except ValueError:
                messagebox.showerror("Invalid", "Enter numeric value.", parent=win)
                return

            output_qty = get_output_qty()

            for item in sel:
                vals = list(tree.item(item, "values"))
                available_qty = float(vals[3])
                total_needed = round(output_qty * qty_per_output, 2)
                missing = round(max(0.0, total_needed - available_qty), 2)
                vals[7] = round(qty_per_output, 2)
                vals[4] = total_needed
                vals[5] = missing
                tree.item(item, values=vals)

        def save_order():
            try:
                output_qty = float(output_qty_entry.get().strip())
                if output_qty <= 0:
                    raise ValueError
            except ValueError:
                messagebox.showerror("Invalid", "Output quantity must be a positive number.", parent=win)
                return

            items = []
            for iid in tree.get_children():
                vals = tree.item(iid, "values")
                mc_id = vals[0]
                total_needed = float(vals[4])
                qty_per_output = float(vals[7])

                if qty_per_output > 0:
                    items.append((mc_id, qty_per_output, total_needed))

            if not items:
                messagebox.showwarning("Empty order", "No material has quantity per output specified.", parent=win)
                return

            order_name = name_entry.get().strip()
            product_name = product_entry.get().strip()
            output_unit = output_unit_combo.get().strip()
            notes = notes_entry.get().strip()

            if not order_name:
                messagebox.showwarning("Missing data", "Enter an order name.", parent=win)
                return

            if not product_name:
                messagebox.showwarning("Missing data", "Enter the produced item.", parent=win)
                return

            conn = get_conn()
            cur = conn.cursor()

            if is_edit:
                cur.execute("""
                    UPDATE orders
                    SET name = ?, product_name = ?, output_qty = ?, output_unit = ?, notes = ?
                    WHERE id = ?
                """, (
                    order_name if order_name else None,
                    product_name if product_name else None,
                    output_qty,
                    output_unit,
                    notes if notes else None,
                    order_id
                ))

                cur.execute("DELETE FROM order_items WHERE order_id = ?", (order_id,))

                current_order_id = order_id
            else:
                cur.execute("""
                    INSERT INTO orders (branch_id, name, product_name, output_qty, output_unit, created_by, created_at, notes)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                """, (
                    self.current_branch_id,
                    order_name if order_name else None,
                    product_name if product_name else None,
                    output_qty,
                    output_unit,
                    self.user["id"],
                    datetime.datetime.utcnow().isoformat(),
                    notes if notes else None
                ))
                current_order_id = cur.lastrowid

            for mc_id, qty_per_output, total_needed in items:
                cur.execute("""
                    INSERT INTO order_items (order_id, material_color_id, qty_per_piece, needed_qty)
                    VALUES (?, ?, ?, ?)
                """, (current_order_id, mc_id, qty_per_output, total_needed))

            conn.commit()
            conn.close()

            messagebox.showinfo("Saved", f"Order '{order_name}' saved successfully.", parent=win)
            win.destroy()
            self.load_orders()
            self.load_materials()

        btns = ttk.Frame(win)
        btns.pack(pady=8)

        ttk.Button(
            btns,
            text="Set qty / output for selection",
            command=set_qty_for_selected
        ).pack(side="left", padx=6)

        ttk.Button(
            btns,
            text="Save Order",
            command=save_order
        ).pack(side="left", padx=6)

        ttk.Button(
            btns,
            text="Cancel",
            command=win.destroy
        ).pack(side="left", padx=6)

    def open_order_items(self, event):
        sel = self.order_tree.selection()
        if not sel:
            return

        values = self.order_tree.item(sel[0])["values"]
        order_id = values[0]
        order_name = values[1]

        win = tk.Toplevel(self)
        win.title(f"Order #{order_id} - {order_name or 'No Name'}")
        win.geometry("920x420")
        win.transient(self)
        win.grab_set()
        self.center_toplevel(win)

        cols = ("mc_id", "ref", "color", "needed_qty", "unit", "qty_per_piece", "created_at")
        tree = self.create_tree_with_scrollbars(win, cols)

        specs = [
            ("mc_id", "ID", 70, "center", True),
            ("ref", "Reference", 220, "w", True),
            ("color", "Color", 110, "w", True),
            ("needed_qty", "Total Needed", 100, "center", True),
            ("unit", "Unit", 70, "center", True),
            ("qty_per_piece", "Qty / Output", 100, "center", True),
            ("created_at", "Created At", 170, "w", True),
        ]
        self.setup_columns(tree, specs)
        self.bind_header_autofit(tree)

        conn = get_conn()
        cur = conn.cursor()
        cur.execute("""
            SELECT
                oi.material_color_id,
                m.reference,
                mc.color,
                oi.needed_qty,
                mc.unit,
                oi.qty_per_piece,
                o.created_at
            FROM order_items oi
            JOIN orders o ON oi.order_id = o.id
            JOIN material_colors mc ON oi.material_color_id = mc.id
            JOIN materials m ON mc.material_id = m.id
            WHERE oi.order_id = ?
        """, (order_id,))

        for r in cur.fetchall():
            tree.insert("", "end", values=(
                r[0],
                r[1],
                r[2],
                float(r[3] or 0),
                r[4],
                float(r[5] or 0),
                self.format_datetime(r[6])
            ))

        conn.close()

    def show_materials_summary_report(self):
        if not self.current_branch_id:
            messagebox.showwarning("Choose branch", "Select branch first.", parent=self)
            return

        win = tk.Toplevel(self)
        win.title("Materials Summary Report")
        win.geometry("980x420")
        win.resizable(False, False)
        win.transient(self)
        win.grab_set()
        self.center_toplevel(win)

        cols = ("ref", "color", "available", "needed", "missing", "orders", "unit")
        tree = self.create_tree_with_scrollbars(win, cols)

        specs = [
            ("ref", "Reference", 230, "w", True),
            ("color", "Color", 120, "w", True),
            ("available", "Available", 100, "center", True),
            ("needed", "Needed", 100, "center", True),
            ("missing", "Missing", 100, "center", True),
            ("orders", "Orders", 80, "center", True),
            ("unit", "Unit", 70, "center", True),
        ]
        self.setup_columns(tree, specs)
        self.bind_header_autofit(tree)

        conn = get_conn()
        cur = conn.cursor()
        cur.execute("""
            SELECT
                m.reference,
                mc.color,
                mc.available_qty,
                IFNULL(SUM(oi.needed_qty), 0) AS needed,
                COUNT(DISTINCT oi.order_id) AS orders_count,
                mc.unit
            FROM materials m
            JOIN material_colors mc ON mc.material_id = m.id
            LEFT JOIN order_items oi ON oi.material_color_id = mc.id
            LEFT JOIN orders o ON oi.order_id = o.id AND o.branch_id = m.branch_id
            WHERE m.branch_id = ?
            GROUP BY m.reference, mc.color, mc.available_qty, mc.unit
            ORDER BY m.reference, mc.color
        """, (self.current_branch_id,))

        for ref, color, avail, needed, orders_count, unit in cur.fetchall():
            avail = float(avail or 0)
            needed = float(needed or 0)
            missing = max(0.0, needed - avail)

            tree.insert("", "end", values=(
                ref,
                color,
                avail,
                needed,
                missing,
                int(orders_count or 0),
                unit
            ))
        conn.close()

    def export_materials_summary(self):
        if not self.current_branch_id:
            messagebox.showwarning("Choose branch", "Select branch first.", parent=self)
            return

        file_path = filedialog.asksaveasfilename(
            parent=self,
            defaultextension=".xlsx",
            filetypes=[("Excel Workbook", "*.xlsx")],
            initialfile=f"branch_{self.current_branch_id}_materials_summary.xlsx"
        )
        if not file_path:
            return

        conn = get_conn()
        cur = conn.cursor()
        cur.execute("""
            SELECT
                m.reference,
                mc.color,
                mc.available_qty,
                IFNULL(SUM(oi.needed_qty), 0) AS needed,
                MAX(0, IFNULL(SUM(oi.needed_qty), 0) - mc.available_qty) AS missing,
                COUNT(DISTINCT oi.order_id) AS orders_count,
                mc.unit
            FROM materials m
            JOIN material_colors mc ON mc.material_id = m.id
            LEFT JOIN order_items oi ON oi.material_color_id = mc.id
            LEFT JOIN orders o ON oi.order_id = o.id AND o.branch_id = m.branch_id
            WHERE m.branch_id = ?
            GROUP BY m.reference, mc.color, mc.available_qty, mc.unit
            ORDER BY m.reference, mc.color
        """, (self.current_branch_id,))
        rows = cur.fetchall()
        conn.close()

        wb = Workbook()
        ws = wb.active
        ws.title = "Materials Summary"

        ws.append(["Reference", "Color", "Available", "Needed", "Missing", "Orders", "Unit"])

        for row in rows:
            ws.append(list(row))

        for column_cells in ws.columns:
            max_length = 0
            column_letter = column_cells[0].column_letter
            for cell in column_cells:
                if cell.value is not None:
                    max_length = max(max_length, len(str(cell.value)))
            ws.column_dimensions[column_letter].width = max_length + 2

        wb.save(file_path)
        messagebox.showinfo("Exported", f"Materials summary exported to:\n{file_path}", parent=self)

    def export_orders_summary(self):
        if not self.current_branch_id:
            messagebox.showwarning("Choose branch", "Select branch first.", parent=self)
            return

        file_path = filedialog.asksaveasfilename(
            parent=self,
            defaultextension=".xlsx",
            filetypes=[("Excel Workbook", "*.xlsx")],
            initialfile=f"branch_{self.current_branch_id}_orders_summary.xlsx"
        )
        if not file_path:
            return

        conn = get_conn()
        cur = conn.cursor()
        cur.execute("""
            SELECT
                o.id,
                o.name,
                o.product_name,
                o.output_qty,
                o.output_unit,
                u.username,
                o.created_at,
                o.notes
            FROM orders o
            LEFT JOIN users u ON o.created_by = u.id
            WHERE o.branch_id = ?
            ORDER BY o.created_at DESC
        """, (self.current_branch_id,))
        rows = cur.fetchall()
        conn.close()

        wb = Workbook()
        ws = wb.active
        ws.title = "Orders Summary"

        ws.append(
            ["Order ID", "Order Name", "Produced Item", "Output Qty", "Unit", "Created By", "Created At", "Notes"])

        for row in rows:
            row = list(row)
            row[6] = self.format_datetime(row[6])
            ws.append(row)

        for column_cells in ws.columns:
            max_length = 0
            column_letter = column_cells[0].column_letter
            for cell in column_cells:
                if cell.value is not None:
                    max_length = max(max_length, len(str(cell.value)))
            ws.column_dimensions[column_letter].width = max_length + 2

        wb.save(file_path)
        messagebox.showinfo("Exported", f"Orders summary exported to:\n{file_path}", parent=self)